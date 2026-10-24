from docx import Document
from pathlib import Path
from lxml import etree as E
from zipfile import ZipFile
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import simpleSplit
import json
p=Path('output/SRGOP_Proposta_Tecnologica_Revisada_Etapa_2.docx'); d=Document(p)
text='\n'.join(p.text for p in d.paragraphs)
for marker in ['XXXX','NOME DO ALUNO','Fazer uma breve','Descreva o propósito','Liste ações','Nesta seção']:
 assert marker not in text, marker
assert len(d.sections)==2
assert len(d.tables)==1
assert len(d.tables[0].rows)==9
print('Estrutura: duas seções, tabela com nove linhas, campos do modelo preenchidos.')
pdfmetrics.registerFont(TTFont('Arial','C:/Windows/Fonts/arial.ttf'))
for title,blocks in json.loads(Path('.runtime/relatorio/content.json').read_text(encoding='utf8')):
 height=30
 for kind,t in blocks:
  if kind=='table': height+=240; continue
  lines=simpleSplit(t,'Arial',12,453.5-(35.45 if kind=='p' else 0))
  height+=len(lines)*(12 if kind in ['ref','h'] else 18)+(9 if kind=='ref' else 7 if kind=='h' else 5)
 print(title,round(height),'pt estimados / 700 disponíveis')
