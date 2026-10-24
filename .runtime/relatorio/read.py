from docx import Document
from pypdf import PdfReader
from pathlib import Path
p=PdfReader(r'D:/Downloads/Relatório Técnico MVP.pdf').pages
for i,page in enumerate(p):
 print(f'\n--- PDF PÁGINA {i+1} ---\n'+page.extract_text())
d=Document(r'D:/Downloads/Template – Proposta Tecnológica (1).docx')
print('\n--- TEMPLATE ---')
for i,p in enumerate(d.paragraphs):
 print(i,repr(p.text),p.style.name)
for t in d.tables:
 print('TABLE',[[c.text for c in r.cells] for r in t.rows])
for s in d.sections: print('SECTION',s.page_width.cm,s.page_height.cm,s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm)

